"""Excel without Qt (khervesheet.core.xlsx): read into the shared layout,
write the engine's workbook back."""

import io

import pytest

openpyxl = pytest.importorskip("openpyxl")

from khervesheet.core.engine import Workbook  # noqa: E402
from khervesheet.core.xlsx import (  # noqa: E402
    _is_excel_compatible, number_format_from_xl, read_xlsx, write_xlsx,
)


def _sample() -> bytes:
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Month", "Sales", "Share"])
    for month, sales in [("Jan", 10), ("Feb", 14), ("Mar", 9)]:
        ws.append([month, sales, None])
    for r in range(2, 5):
        ws.cell(r, 3).value = f"=B{r}/SUM($B$2:$B$4)"
        ws.cell(r, 3).number_format = "0.00%"
    ws["A1"].font = Font(bold=True, color="FFFF0000")
    ws["B1"].fill = PatternFill("solid", start_color="FFFFFF00")
    ws["B2"].alignment = Alignment(horizontal="center")
    ws["B3"].border = Border(bottom=Side(style="medium"))
    ws.column_dimensions["A"].width = 20
    ws.freeze_panes = "B2"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    chart.title = "Sales"
    ws.add_chart(chart, "E2")
    wb.create_sheet("Other")["A1"] = "=Data!B2*2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read():
    layout = read_xlsx(_sample())
    data, other = layout["sheets"]
    assert data["name"] == "Data" and data["rows"] == 5000
    assert [1, 2, "=B2/SUM($B$2:$B$4)"] in data["cells"]
    assert [1, 1, "10"] in data["cells"]
    formats = {(r, c): f for r, c, f in data["formats"]}
    assert formats[(0, 0)] == {"bold": True, "font_color": "#FF0000"}
    assert formats[(0, 1)] == {"bg": "#FFFF00"}
    assert formats[(1, 1)]["alignment"] == 0x84
    assert formats[(2, 1)]["b_bottom"]["width"] == 2.0
    assert formats[(1, 2)]["number_format"] == "Percentage"
    assert data["widths"] == [[0, 145]]
    assert (data["freezeRows"], data["freezeCols"]) == (1, 1)
    assert other["cells"] == [[0, 0, "=Data!B2*2"]]
    (chart,) = layout["charts"]
    assert chart["type"] == "Bar" and chart["title"] == "Sales"
    assert chart["x"] == "A2:A4" and chart["series"][0]["ref"] == "B2:B4"
    assert chart["series"][0]["name"] == "Sales"
    assert (chart["row"], chart["col"]) == (1, 4)


def test_write_round_trip():
    layout = read_xlsx(_sample())
    engine = Workbook()
    for s in layout["sheets"]:
        engine.add_sheet(s["name"])
    for s in layout["sheets"]:
        sheet = engine.sheet_by_name(s["name"])
        for r, c, src in sorted(s["cells"], key=lambda x: x[2][:1] == "="):
            sheet.set_cell(r, c, src)
    engine.sheet_by_name("Data").set_cell(6, 0, "=PY 1+1")
    engine.sheet_by_name("Data").set_cell(7, 0, "=CUSTOMFN(1)")
    out = write_xlsx(engine, {"sheets": {s["name"]: s
                                         for s in layout["sheets"]},
                              "charts": layout["charts"]})
    back = openpyxl.load_workbook(io.BytesIO(out))
    ws = back["Data"]
    assert ws["C2"].value == "=B2/SUM($B$2:$B$4)"
    assert ws["C2"].number_format == "0.00%"
    assert ws["B2"].value == 10 and isinstance(ws["B2"].value, int)
    assert ws["A1"].font.b and ws["B1"].fill.fgColor.rgb == "FFFFFF00"
    assert ws["B3"].border.bottom.style == "medium"
    assert ws.freeze_panes == "B2"
    assert ws.column_dimensions["A"].width == pytest.approx(20)
    assert len(ws._charts) == 1
    assert back["Other"]["A1"].value == "=Data!B2*2"
    # Python and unknown functions: their shown value, not the code.
    assert not str(ws["A7"].value or "").startswith("=")
    assert not str(ws["A8"].value or "").startswith("=")


def test_number_formats_and_formulas():
    assert number_format_from_xl("0.00") == "0.00"
    assert number_format_from_xl("#,##0") == "0"
    assert number_format_from_xl("yyyy-mm-dd") == "Date:%Y-%m-%d"
    assert number_format_from_xl("dd/mm/yyyy") == "Date:%d/%m/%Y"
    assert number_format_from_xl("h:mm") == "Time:%H:%M"
    assert number_format_from_xl('[$$-409]#,##0.00') == "Currency:$"
    assert number_format_from_xl("0.00E+00") == "Scientific"
    assert number_format_from_xl("General") is None
    assert _is_excel_compatible("=SUM(A1:A3)*2")
    assert not _is_excel_compatible("=PY\nimport numpy")
